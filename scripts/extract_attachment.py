#!/usr/bin/env python3

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


MAX_CHARS = 8000
MAX_ROWS = 25
MAX_COLS = 10
MAX_PDF_PAGES = 6


def normalize_text(text: str) -> str:
    compact = " ".join(text.split())
    if len(compact) > MAX_CHARS:
        return compact[: MAX_CHARS - 3] + "..."
    return compact


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    texts = []

    for page_index, page in enumerate(reader.pages[:MAX_PDF_PAGES]):
        page_text = page.extract_text() or ""
        if page_text.strip():
            texts.append(f"[page {page_index + 1}] {page_text}")

    merged = normalize_text("\n".join(texts))
    return {
        "kind": "document",
        "summary": f"Loaded PDF attachment {path}",
        "text": merged or f"PDF attachment {path} contained no extractable text.",
    }


def extract_delimited(path: Path, delimiter: str) -> dict:
    rows = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row_index, row in enumerate(reader):
            if row_index >= MAX_ROWS:
                break
            rows.append("\t".join(row[:MAX_COLS]))

    merged = normalize_text("\n".join(rows))
    return {
        "kind": "table",
        "summary": f"Loaded tabular attachment {path}",
        "text": merged or f"Tabular attachment {path} contained no rows.",
    }


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet_chunks = []

    for sheet in workbook.worksheets[:3]:
        sheet_chunks.append(f"[sheet] {sheet.title}")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index >= MAX_ROWS:
                break
            values = ["" if value is None else str(value) for value in row[:MAX_COLS]]
            sheet_chunks.append("\t".join(values))

    merged = normalize_text("\n".join(sheet_chunks))
    return {
        "kind": "table",
        "summary": f"Loaded spreadsheet attachment {path}",
        "text": merged or f"Spreadsheet attachment {path} contained no visible cells.",
    }


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: extract_attachment.py <path>", file=sys.stderr)
        return 1

    path = Path(sys.argv[1]).resolve()
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        result = extract_pdf(path)
    elif suffix == ".csv":
        result = extract_delimited(path, ",")
    elif suffix == ".tsv":
        result = extract_delimited(path, "\t")
    elif suffix in {".xlsx", ".xlsm"}:
        result = extract_xlsx(path)
    else:
        print(f"Unsupported attachment extension: {suffix}", file=sys.stderr)
        return 1

    print(json.dumps(result, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
